# -*- coding: utf-8 -*-
"""Content-level verification of the final docx."""
import re
import sys
import zipfile

from docx import Document

path = sys.argv[1]
doc = Document(path)

paras = [p.text for p in doc.paragraphs]
tbl_texts = []
for t in doc.tables:
    for row in t.rows:
        for c in row.cells:
            tbl_texts.append(c.text)
alltext = "\n".join(paras + tbl_texts)

z = zipfile.ZipFile(path)
xml = z.read("word/document.xml").decode("utf-8")
print("vertAlign subscript:", xml.count('w:val="subscript"'))
print("vertAlign superscript:", xml.count('w:val="superscript"'))

checks = [
    # de-AIGC markers
    ("Disasters such as earthquakes, floods, or landslides frequently damage", True),
    ("well suited to this setting", True),
    ("Three gaps remain.", True),
    ("Shannon and Weaver [7], [8] distinguished", True),
    ("did not make the learned agents winners on raw recall", True),
    ("Future work has four directions", True),
    ("Learning is preceded by an explicit warm-up", True),
    ("so the comparison is paired rather than marginal", True),
    # things that must be GONE
    ("This paper makes the following contributions", False),
    ("The remainder of the paper is organized as follows", False),
    ("natural fit because", False),
    ("Summarizing the above literature", False),
    ("In short, the benchmark measures", False),
    ("did not crown the learned agents", False),
    ("genuine advantage is efficiency", False),
    ("markedly", False),
    ("comprehensive tutorial", False),
    ("extensively studied", False),
    ("First, the channel, utility, and recall models are parametric", False),
    ("## VI", False),
    # consolidation check: claim appears exactly twice (VI-D opening + VII)
    ("The correction removes the zero-recall failure mode but does not make the learned agents winners on raw recall", "==2"),
    # honest framing intact
    ("0.346", True),
    ("40% of the bit cost", True),
    ("288.0 bits per step", True),
    ("economically irrational", True),
    ("we state this as a hypothesis rather than a conclusion", True),
    ("A tentative interpretation", True),
    ("could act as a confounding factor", True),
    # structure
    ("Recall-aware reward weight introduced in the corrected design", True),
    ("[16]", True),
    ("Eq.", False),
]
fails = 0
for s, want in checks:
    n = alltext.count(s)
    if want == "==2":
        ok = n == 2
    else:
        ok = (n > 0) if want else (n == 0)
    if not ok:
        fails += 1
    print("%s  n=%d  %s" % ("OK " if ok else "FAIL", n, s[:70]))
print("reference items in doc:", sum(1 for p in paras if re.match(r"^\[\d+\]", p)))
print("FAILS:", fails)
